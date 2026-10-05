import io
import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF

def sanitize_text(text: str) -> str:
    """Removes unsupported typographic characters for clean PDF/Word exports."""
    replacements = {
        '“': '"', '”': '"', "‘": "'", "’": "'",
        '—': '-', '–': '-', '…': '...'
    }
    for orig, repl in replacements.items():
        text = text.replace(orig, repl)
    return text.strip()

def format_html_preview(text: str) -> str:
    """Converts raw text into styled HTML for rendering inside the web browser."""
    lines = text.split('\n')
    html_out = []
    for line in lines:
        line_s = line.strip()
        if line_s.startswith("## "):
            html_out.append(f"<h3 style='color: #4CAF50; margin-top: 15px;'>{line_s[3:]}</h3>")
        elif line_s.startswith("# "):
            html_out.append(f"<h2 style='color: #2196F3; margin-top: 20px;'>{line_s[2:]}</h2>")
        elif line_s.startswith("**") and line_s.endswith("**"):
            html_out.append(f"<strong>{line_s[2:-2]}</strong><br/>")
        elif line_s == "":
            html_out.append("<br/>")
        else:
            html_out.append(f"<p style='margin: 4px 0; line-height: 1.5;'>{line_s}</p>")
    return "".join(html_out)

def format_docx(text: str, doc_type: str, logo_path: str = None) -> io.BytesIO:
    """Creates a formatted Word document (.docx) with title, text, and optional logo."""
    doc = Document()
    
    if logo_path and os.path.exists(logo_path):
        header_para = doc.add_paragraph()
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = header_para.add_run()
        run.add_picture(logo_path, width=Inches(1.8))

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(doc_type.upper())
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.name = "Times New Roman"

    for paragraph in text.split('\n'):
        p = doc.add_paragraph()
        run = p.add_run(paragraph)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

class LegalPDF(FPDF):
    def __init__(self, logo_path=None):
        super().__init__()
        self.logo_path = logo_path

    def header(self):
        if self.logo_path and os.path.exists(self.logo_path):
            self.image(self.logo_path, 10, 8, 33)
            self.ln(15)

    def footer(self):
        self.set_y(-15)
        self.set_font('helvetica', 'I', 8)
        self.set_text_color(128)
        self.cell(0, 10, 'LegalEase Inc. | contact@legalease.com | All Rights Reserved.', align='C')

def format_pdf(text: str, doc_type: str, logo_path: str = None) -> bytes:
    """Creates a formatted PDF document (.pdf) with standard legal fonts and layout."""
    pdf = LegalPDF(logo_path=logo_path)
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    
    pdf.set_font("helvetica", "B", 16)
    pdf.cell(0, 10, sanitize_text(doc_type.upper()), new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(5)

    pdf.set_font("helvetica", size=10)
    sanitized = sanitize_text(text)
    for line in sanitized.split('\n'):
        if line.startswith("## ") or line.startswith("# "):
            pdf.set_font("helvetica", "B", 12)
            pdf.multi_cell(0, 7, line.replace("#", "").strip())
            pdf.set_font("helvetica", size=10)
        else:
            pdf.multi_cell(0, 5, line)
            pdf.ln(2)

    return bytes(pdf.output())